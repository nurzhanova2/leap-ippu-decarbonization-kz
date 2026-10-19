"""Inventory versions shared by the CRT extraction and analysis scripts.

crt2025 is the CRT 2025 submission (1990-2023) used by the conference
version; its outputs keep their original file names and are archived.
crt2026 is the CRT 2026 submission (1990-2024) used by the journal version;
its outputs carry the suffix "_crt2026".

Scripts select a version with "--inventory crt2025|crt2026" (default crt2025).
CRT 2026 workbooks are read directly from the submitted ZIP, so nothing is
extracted into data/raw/.
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


VERSIONS = {
    "crt2025": {
        "source": "UNFCCC CRT 2025",
        "submission_date": "2025-04-15",
        "path": Path("data/raw/unfccc_crt_2025/extracted"),
        "last_year": 2023,
        "suffix": "",
    },
    "crt2026": {
        "source": "UNFCCC CRT 2026",
        "submission_date": "2026-04-16",
        "path": Path("data/raw/unfccc_crt_2026/KAZ-CRT-2026-V1.0.zip"),
        "last_year": 2024,
        "suffix": "_crt2026",
    },
}


def selected_version(default: str = "crt2025") -> str:
    """Return the version named by --inventory on the command line."""
    if "--inventory" in sys.argv:
        version = sys.argv[sys.argv.index("--inventory") + 1]
        if version not in VERSIONS:
            raise SystemExit(f"Unknown inventory {version!r}; choose from {', '.join(VERSIONS)}")
        return version
    return default


def workbook_year(name: str) -> str:
    """Inventory year from a CRT file name (2025 and 2026 naming)."""
    return re.search(r"V1\.0-(\d{4})(?:-|\.xlsx$)", name).group(1)


def crt_workbooks(version: str):
    """Yield (file name, workbook) for every country-year CRT workbook, sorted by name."""
    path = VERSIONS[version]["path"]
    if path.suffix == ".zip":
        with zipfile.ZipFile(path) as archive:
            for name in sorted(n for n in archive.namelist() if n.endswith(".xlsx")):
                yield name, load_workbook(io.BytesIO(archive.read(name)), read_only=True, data_only=True)
    else:
        for workbook in sorted(path.glob("*.xlsx")):
            yield workbook.name, load_workbook(workbook, read_only=True, data_only=True)
