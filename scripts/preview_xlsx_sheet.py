"""Print non-empty cells from the first rows and columns of an XLSX sheet."""

import argparse
from openpyxl import load_workbook


parser = argparse.ArgumentParser()
parser.add_argument("workbook")
parser.add_argument("sheet")
parser.add_argument("--rows", type=int, default=35)
parser.add_argument("--cols", type=int, default=16)
args = parser.parse_args()

book = load_workbook(args.workbook, read_only=True, data_only=False)
worksheet = book[args.sheet]
for row_number, row in enumerate(worksheet.iter_rows(max_row=args.rows, max_col=args.cols, values_only=True), start=1):
    values = [(index + 1, value) for index, value in enumerate(row) if value is not None]
    if values:
        print(row_number, values)
