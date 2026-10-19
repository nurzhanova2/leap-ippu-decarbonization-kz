"""Extract IPPU activity data reported in CRT and CRF 2023 Table 2(I).A-H.

--inventory crt2025 (default) writes CRF 2023 and CRT 2025 rows to the archived
file; --inventory crt2026 writes CRT 2026 rows to a file with the suffix _crt2026.
"""

import csv
from pathlib import Path
import re
from openpyxl import load_workbook

from inventory_versions import VERSIONS, crt_workbooks, selected_version, workbook_year


CRF_DIR = Path("data/raw/unfccc_crf_2023/extracted")
VERSION = selected_version()
OUTPUT = Path(f"data/processed/unfccc_ippu_activity_data_1990_{VERSIONS[VERSION]['last_year']}{VERSIONS[VERSION]['suffix']}.csv")


def numeric(value):
    return float(value) if isinstance(value, (int, float)) else None


def crt_rows(name, book, version):
    sheet = book["Table2(I).A-H"]
    records = []
    for row in sheet.iter_rows(values_only=True):
        label = row[1] if len(row) > 1 else None
        activity = numeric(row[3] if len(row) > 3 else None)
        if not isinstance(label, str) or activity is None:
            continue
        match = re.match(r"^(2(?:\.[A-H](?:\.\d+(?:\.[a-z]+)?)?)?)\.\s*(.*)$", label.strip())
        if not match:
            continue
        records.append({
            "source": VERSIONS[version]["source"],
            "submission_date": VERSIONS[version]["submission_date"],
            "inventory_year": workbook_year(name),
            "category_code": match.group(1),
            "category_label": match.group(2).strip(),
            "activity_description": row[2],
            "activity_value_kt": activity,
            "activity_value_tonnes": activity * 1000,
            "implied_co2_ef_t_per_t": numeric(row[4] if len(row) > 4 else None),
            "source_table": "Table2(I).A-H",
        })
    return records


def crf_rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["Table2(I).A-Hs1"]
    parent = None
    records = []
    for row in sheet.iter_rows(values_only=True):
        raw_label = row[0]
        activity = numeric(row[2] if len(row) > 2 else None)
        if not isinstance(raw_label, str) or activity is None:
            continue
        label = raw_label.strip()
        if match := re.match(r"^([A-H])\.\s*(.*)$", label):
            parent, code, short_label = match.group(1), f"2.{match.group(1)}", match.group(2).strip()
        elif match := re.match(r"^(\d+)\.\s*(.*)$", label):
            if parent is None:
                continue
            code, short_label = f"2.{parent}.{match.group(1)}", match.group(2).strip()
        elif match := re.match(r"^([a-z])\.\s*(.*)$", label):
            if parent is None:
                continue
            code, short_label = f"2.{parent}.other.{match.group(1)}", match.group(2).strip()
        else:
            continue
        records.append({
            "source": "UNFCCC CRF 2023",
            "submission_date": "2023-04-15",
            "inventory_year": re.search(r"KAZ_2023_(\d{4})_", path.name).group(1),
            "category_code": code,
            "category_label": short_label,
            "activity_description": row[1],
            "activity_value_kt": activity,
            "activity_value_tonnes": activity * 1000,
            "implied_co2_ef_t_per_t": numeric(row[3] if len(row) > 3 else None),
            "source_table": "Table2(I).A-Hs1",
        })
    return records


records = []
if VERSION == "crt2025":
    for workbook in sorted(CRF_DIR.glob("*.xlsx")):
        records.extend(crf_rows(workbook))
for name, book in crt_workbooks(VERSION):
    records.extend(crt_rows(name, book, VERSION))

with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=records[0].keys())
    writer.writeheader()
    writer.writerows(records)
print(f"Wrote {len(records)} records to {OUTPUT}")
